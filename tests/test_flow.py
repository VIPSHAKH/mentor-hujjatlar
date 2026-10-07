import base64
import io
import zipfile

from openpyxl import Workbook, load_workbook

from tests.conftest import PDF, PNG, district_id, make_mentor


def _doc(me: dict, code: str) -> dict:
    return next(d for d in me["documents"] if d["type"]["code"] == code)


async def test_full_review_cycle(client, admin_headers):
    mentor, mh = await make_mentor(client, admin_headers, "Aliyev Anvar Anvarovich", "+998901112233")

    me = (await client.get("/api/mentor/me", headers=mh)).json()
    assert [d["type"]["code"] for d in me["documents"]] == ["passport", "diploma", "certificate", "objectivka"]
    assert all(d["status"] == "MISSING" for d in me["documents"])
    assert me["overall"] == "MISSING"

    r = await client.post("/api/mentor/documents/diploma", headers=mh, files={"file": ("diplom.pdf", PDF, "application/pdf")})
    assert r.status_code == 200, r.text
    assert "muvaffaqiyatli yuklandi" in r.json()["message"]

    # Admin bildirishnoma oladi
    n = (await client.get("/api/admin/notifications", headers=admin_headers)).json()
    assert n["unread"] >= 1 and "Aliyev" in n["items"][0]["title"]

    queue = (await client.get("/api/admin/review-queue", headers=admin_headers)).json()
    item = next(i for i in queue if i["mentor"]["id"] == mentor["id"])
    v1 = item["version_id"]

    # Sababsiz qaytarib bo'lmaydi
    r = await client.post(f"/api/admin/versions/{v1}/return", headers=admin_headers, json={"comment": ""})
    assert r.status_code == 422
    r = await client.post(
        f"/api/admin/versions/{v1}/return",
        headers=admin_headers,
        json={"comment": "Diplomning to'liq nusxasi yuklanmagan."},
    )
    assert r.status_code == 200

    me = (await client.get("/api/mentor/me", headers=mh)).json()
    d = _doc(me, "diploma")
    assert d["status"] == "RETURNED"
    assert d["return_comment"] == "Diplomning to'liq nusxasi yuklanmagan."
    assert me["overall"] == "RETURNED"
    notes = (await client.get("/api/mentor/notifications", headers=mh)).json()
    assert any(x["kind"] == "returned" and "Sabab:" in x["body"] for x in notes)

    # Qayta yuklash: 2-versiya, eski fayl tarixda qoladi
    r = await client.post("/api/mentor/documents/diploma", headers=mh, files={"file": ("diplom2.png", PNG, "image/png")})
    assert r.status_code == 200
    detail = (await client.get(f"/api/admin/mentors/{mentor['id']}", headers=admin_headers)).json()
    d = next(x for x in detail["documents"] if x["type"]["code"] == "diploma")
    assert [v["version_no"] for v in d["versions"]] == [2, 1]
    assert [v["status"] for v in d["versions"]] == ["PENDING", "RETURNED"]
    v2 = d["versions"][0]["id"]

    # Ikkala versiya fayli ham ochiladi
    r = await client.get(f"/api/admin/files/{v1}", headers=admin_headers)
    assert r.status_code == 200 and r.content == PDF and r.headers["content-type"] == "application/pdf"
    r = await client.get(f"/api/admin/files/{v2}?download=1", headers=admin_headers)
    assert r.content == PNG and r.headers["content-disposition"].startswith("attachment")

    assert (await client.post(f"/api/admin/versions/{v2}/approve", headers=admin_headers)).status_code == 200
    # Tasdiqlangan hujjatni qayta yuklab bo'lmaydi
    r = await client.post("/api/mentor/documents/diploma", headers=mh, files={"file": ("x.pdf", PDF, "application/pdf")})
    assert r.status_code == 409

    for code in ("passport", "certificate", "objectivka"):
        await client.post(f"/api/mentor/documents/{code}", headers=mh, files={"file": (f"{code}.pdf", PDF, "application/pdf")})
    detail = (await client.get(f"/api/admin/mentors/{mentor['id']}", headers=admin_headers)).json()
    assert detail["overall"] == "PENDING"
    for d in detail["documents"]:
        if d["status"] == "PENDING":
            await client.post(f"/api/admin/versions/{d['latest']['id']}/approve", headers=admin_headers)
    me = (await client.get("/api/mentor/me", headers=mh)).json()
    assert me["overall"] == "COMPLETE"
    notes = (await client.get("/api/mentor/notifications", headers=mh)).json()
    assert any(x["kind"] == "all_approved" for x in notes)

    # ZIP papka
    r = await client.get(f"/api/admin/mentors/{mentor['id']}/archive.zip", headers=admin_headers)
    names = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert "Gurlan tumani/Aliyev Anvar Anvarovich/Diplom.png" in names and len(names) == 4

    # Amallar tarixi
    log = (await client.get(f"/api/admin/mentors/{mentor['id']}/audit", headers=admin_headers)).json()
    actions = {a["action"] for a in log}
    assert {"document.upload", "document.return", "document.approve", "mentor.complete"} <= actions


async def test_upload_validation(client, admin_headers):
    _, mh = await make_mentor(client, admin_headers, "Validatsiya Test", "+998901112244")
    r = await client.post(
        "/api/mentor/documents/passport", headers=mh, files={"file": ("a.pdf", b"MZ\x90\x00evil", "application/pdf")}
    )
    assert r.status_code == 400
    r = await client.post(
        "/api/mentor/documents/passport", headers=mh, files={"file": ("big.pdf", PDF + b"0" * (1024 * 1024), "application/pdf")}
    )
    assert r.status_code == 413
    r = await client.post("/api/mentor/documents/unknown", headers=mh, files={"file": ("a.pdf", PDF, "application/pdf")})
    assert r.status_code == 404


async def test_mentor_isolation(client, admin_headers):
    _, a = await make_mentor(client, admin_headers, "Birinchi Mentor", "+998901112255")
    other, b = await make_mentor(client, admin_headers, "Ikkinchi Mentor", "+998901112266")
    await client.post("/api/mentor/documents/passport", headers=b, files={"file": ("p.pdf", PDF, "application/pdf")})
    detail = (await client.get(f"/api/admin/mentors/{other['id']}", headers=admin_headers)).json()
    vid = detail["documents"][0]["latest"]["id"]
    assert (await client.get(f"/api/mentor/files/{vid}", headers=b)).status_code == 200
    assert (await client.get(f"/api/mentor/files/{vid}", headers=a)).status_code == 404
    # Mentor tokeni admin API'ga kira olmaydi, tokensiz fayl ochilmaydi
    assert (await client.get("/api/admin/mentors", headers=a)).status_code == 401
    assert (await client.get(f"/api/admin/files/{vid}")).status_code == 401


async def test_district_admin_scope(client, admin_headers):
    shovot = await district_id(client, admin_headers, "Shovot tumani")
    r = await client.post(
        "/api/admin/admins",
        headers=admin_headers,
        json={"username": "shovot", "password": "secret12", "role": "DISTRICT", "district_id": shovot},
    )
    assert r.status_code == 200, r.text
    token = (await client.post("/api/admin/login", json={"login": "shovot", "password": "secret12"})).json()["token"]
    h = {"Authorization": f"Bearer {token}"}
    sm, _ = await make_mentor(client, admin_headers, "Shovotlik Mentor", "+998901112277", "Shovot tumani")
    items = (await client.get("/api/admin/mentors", headers=h)).json()["items"]
    assert items and all(i["district"]["name"] == "Shovot tumani" for i in items)
    gurlan = next(
        i
        for i in (await client.get("/api/admin/mentors", headers=admin_headers)).json()["items"]
        if i["district"]["name"] == "Gurlan tumani"
    )
    assert (await client.get(f"/api/admin/mentors/{gurlan['id']}", headers=h)).status_code == 404
    assert (await client.get(f"/api/admin/mentors/{sm['id']}", headers=h)).status_code == 200
    assert len((await client.get("/api/admin/districts", headers=h)).json()) == 1


async def test_filters_stats_export(client, admin_headers):
    r = await client.get("/api/admin/mentors", headers=admin_headers, params={"doc_type": "diploma", "doc_status": "MISSING"})
    names = {i["full_name"] for i in r.json()["items"]}
    assert "Aliyev Anvar Anvarovich" not in names and "Birinchi Mentor" in names

    r = await client.get("/api/admin/mentors", headers=admin_headers, params={"overall": "COMPLETE"})
    assert [i["full_name"] for i in r.json()["items"]] == ["Aliyev Anvar Anvarovich"]

    r = await client.get("/api/admin/mentors", headers=admin_headers, params={"q": "ikkinchi"})
    assert [i["full_name"] for i in r.json()["items"]] == ["Ikkinchi Mentor"]

    dash = (await client.get("/api/admin/dashboard", headers=admin_headers)).json()
    gurlan = next(d for d in dash["districts"] if d["name"] == "Gurlan tumani")
    assert gurlan["total"] == 4 and gurlan["complete"] == 1 and gurlan["incomplete"] == 3
    assert dash["totals"]["total"] == sum(d["total"] for d in dash["districts"])

    r = await client.get("/api/admin/mentors/export.xlsx", headers=admin_headers, params={"district_id": gurlan["id"]})
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.values)
    assert rows[0] == (
        "№",
        "Tuman",
        "F.I.Sh.",
        "Telefon",
        "Pasport yoki ID karta holati",
        "Diplom holati",
        "Xalqaro sertifikat holati",
        "Obektivka holati",
        "Umumiy holat",
    )
    aliyev = next(r for r in rows if r[2] == "Aliyev Anvar Anvarovich")
    assert aliyev[4:] == ("Tasdiqlangan",) * 4 + ("To'liq tasdiqlangan",)


def _xlsx(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_excel_import_grouped_and_tabular(client, admin_headers):
    grouped = _xlsx(
        [
            ["Gurlan tumani"],
            [1, "Rasulov Bekzod"],
            [2, "Karimova Dilnoza Bahodir qizi"],
            ["Yangi tumani"],
            ["Toshmatov Sardor"],
        ]
    )
    r = await client.post("/api/admin/import", headers=admin_headers, files={"file": ("m.xlsx", grouped)})
    body = r.json()
    assert r.status_code == 200, body
    assert body["created"] == 3 and body["new_districts"] == ["Yangi tumani"] and not body["errors"]
    creds = load_workbook(io.BytesIO(base64.b64decode(body["credentials_xlsx"]))).active
    login, password = list(creds.values)[1][2:4]
    assert (await client.post("/api/mentor/login", json={"login": login, "password": password})).status_code == 200

    tabular = _xlsx(
        [
            ["Tuman", "F.I.Sh.", "Tug‘ilgan sana", "Telefon", "JSHSHIR", "Yo‘nalish"],
            ["Gurlan tumani", "Rasulov Bekzod", "01.02.1991", "90 123 45 67", "12345678901234", "Fizika"],
            ["Shovot", "Yangi Mentor", None, "+998 91 765 43 21", None, None],
            ["Shovot tumani", "Xato Telefon", None, "12345", None, None],
        ]
    )
    r = await client.post("/api/admin/import", headers=admin_headers, files={"file": ("m.xlsx", tabular)})
    body = r.json()
    assert body["created"] == 1 and body["updated"] == 1 and len(body["errors"]) == 1
    items = (await client.get("/api/admin/mentors", headers=admin_headers, params={"q": "Rasulov"})).json()["items"]
    assert len(items) == 1 and items[0]["phone"] == "+998901234567" and items[0]["direction"] == "Fizika"
    shovot = (await client.get("/api/admin/mentors", headers=admin_headers, params={"q": "Yangi Mentor"})).json()["items"]
    assert shovot[0]["district"]["name"] == "Shovot tumani"


async def test_new_document_type_backfills(client, admin_headers):
    r = await client.post(
        "/api/admin/document-types",
        headers=admin_headers,
        json={"code": "cv", "name": "CV/rezume", "is_required": False, "sort_order": 10},
    )
    assert r.status_code == 200
    items = (await client.get("/api/admin/mentors", headers=admin_headers, params={"overall": "COMPLETE"})).json()["items"]
    assert items[0]["docs"]["cv"] == "MISSING"  # ixtiyoriy hujjat umumiy holatga ta'sir qilmaydi
    await client.patch(f"/api/admin/document-types/{r.json()['id']}", headers=admin_headers, json={"is_active": False})


async def test_login_throttle(client):
    for _ in range(10):
        r = await client.post("/api/mentor/login", json={"login": "+998999999999", "password": "x"})
        assert r.status_code == 401
    r = await client.post("/api/mentor/login", json={"login": "+998999999999", "password": "x"})
    assert r.status_code == 429
