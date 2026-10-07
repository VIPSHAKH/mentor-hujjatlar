"""Excel import (mentorlar ro'yxati) va eksport."""

import io
import re
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Ustun sarlavhalarini tanish (kichik harf, apostrofsiz)
HEADER_ALIASES = {
    "district": ["tuman", "hudud", "shahar", "tuman/shahar", "tumanshahar"],
    "full_name": ["fish", "fio", "fullname", "ism", "ismfamiliya", "familiyaismsharifi", "mentor"],
    "birth_date": ["tugilgansana", "tugilgankun", "birthdate"],
    "phone": ["telefon", "telefonraqami", "tel", "phone"],
    "passport": ["pasport", "pasportseriya", "pasportseriyasivaraqami", "passport", "pasportid"],
    "pinfl": ["jshshir", "pinfl", "jshshr"],
    "direction": ["yonalish", "yonalishi", "fan", "direction"],
    "activity": ["faoliyat", "faoliyati", "mentorsifatidagifaoliyati", "lavozim"],
    "login": ["login"],
}


def _norm(s) -> str:
    return re.sub(
        r"[^a-z0-9/]", "", str(s or "").lower().replace("‘", "").replace("’", "").replace("'", "").replace("ʻ", "")
    ).replace("/", "")


def _match_header(cell) -> str | None:
    n = _norm(cell)
    if not n:
        return None
    for key, aliases in HEADER_ALIASES.items():
        if n in aliases or any(n.startswith(a) for a in aliases if len(a) > 3):
            return key
    return None


DISTRICT_RE = re.compile(r"(tumani|shahri|tuman|shahar)\b", re.IGNORECASE)


@dataclass
class ParsedRow:
    row: int
    data: dict


@dataclass
class ParseResult:
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def parse_mentors_xlsx(content: bytes) -> ParseResult:
    """Ikki formatni qabul qiladi:
    1) sarlavhali jadval (Tuman, F.I.Sh., Telefon, ...);
    2) guruhlangan ro'yxat: "Gurlan tumani" qatori, ostida F.I.Sh.lar.
    Ikkalasida ham "Tuman" ustuni bo'lmasa, guruh sarlavhasi qatorlari ishlatiladi."""
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        res = ParseResult()
        res.errors.append("Excel faylni o'qib bo'lmadi (.xlsx formatda bo'lishi kerak)")
        return res
    res = ParseResult()
    for ws in wb.worksheets:
        columns: dict[int, str] | None = None
        current_district: str | None = None
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = ["" if c is None else c for c in row]
            non_empty = [(i, c) for i, c in enumerate(cells) if str(c).strip()]
            if not non_empty:
                continue
            if columns is None:
                found = {i: k for i, c in non_empty if (k := _match_header(c))}
                if "full_name" in found.values() and len(found) >= 2:
                    columns = found
                    continue
            # Guruh sarlavhasi: "Gurlan tumani" (bitta katak yoki raqamsiz qator)
            texts = [str(c).strip() for _, c in non_empty if not isinstance(c, (int, float))]
            if len(non_empty) == 1 and texts and DISTRICT_RE.search(texts[0]) and len(texts[0].split()) <= 4:
                current_district = texts[0]
                continue
            if columns is None:
                # Sarlavhasiz ro'yxat: birinchi matnli katak — F.I.Sh. (tartib raqami bo'lsa o'tkaziladi)
                name = next((str(c).strip() for _, c in non_empty if isinstance(c, str) and not c.strip().isdigit()), None)
                if not name or name in ("...", "…"):
                    continue
                data = {"full_name": name, "district": current_district}
            else:
                data = {key: cells[i] if i < len(cells) else None for i, key in columns.items()}
                if not str(data.get("full_name") or "").strip():
                    continue
                if not str(data.get("district") or "").strip():
                    data["district"] = current_district
            data["full_name"] = re.sub(r"\s+", " ", str(data["full_name"]).strip())
            if not data.get("district"):
                res.errors.append(f"{ws.title}!{idx}: «{data['full_name']}» uchun tuman ko'rsatilmagan")
                continue
            data["district"] = re.sub(r"\s+", " ", str(data["district"]).strip())
            res.rows.append(ParsedRow(idx, data))
    return res


STATUS_FILL = {
    "MISSING": "FFF4CC",
    "PENDING": "DCEBFF",
    "APPROVED": "D9F2DC",
    "RETURNED": "FBD5D5",
    "COMPLETE": "D9F2DC",
}


def build_workbook(title: str, headers: list[str], rows: list[list], status_cols: dict[int, list[str]] | None = None) -> bytes:
    """status_cols: ustun indeksi -> har bir qator uchun status kodi (rang berish uchun)."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2A78D6")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for col, codes in (status_cols or {}).items():
        for i, code in enumerate(codes, start=2):
            if color := STATUS_FILL.get(code):
                ws.cell(row=i, column=col + 1).fill = PatternFill("solid", fgColor=color)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(8, width + 2), 50)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


IMPORT_TEMPLATE_HEADERS = [
    "Tuman",
    "F.I.Sh.",
    "Tug‘ilgan sana",
    "Telefon",
    "Pasport",
    "JSHSHIR",
    "Yo‘nalish",
    "Faoliyati",
]


def import_template() -> bytes:
    return build_workbook(
        "Mentorlar",
        IMPORT_TEMPLATE_HEADERS,
        [
            [
                "Gurlan tumani",
                "Aliyev Anvar Anvarovich",
                "15.03.1990",
                "+998901234567",
                "AB1234567",
                "31503900000000",
                "Matematika",
                "Maktab mentori",
            ],
            ["Shovot tumani", "Karimova Dilnoza Bahodir qizi", "02.11.1994", "+998911234567", "", "", "Informatika", ""],
        ],
    )
